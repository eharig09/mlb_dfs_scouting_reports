"""Slate assembly and the workbook it becomes.

Named `test_nfl_*` and carrying no conftest, per the collision documented in
tests/nfl/nfl_fixtures.py.

The join layer's job is to attach sources without inventing anything, so most of what is
checked here is *absence* behaving correctly: a player no source covered must stay blank
through the frame, the figures and the sheet, never becoming a zero that reads as "projected
to score nothing".
"""

import os

import numpy as np
import pandas as pd
import pytest

from nfl import report, report_visuals, slate


def _dk(rows):
    return pd.DataFrame([{
        "Position": r.get("slot", "WR/FLEX"),
        "Name + ID": f"{r['name']} (1)",
        "Name": r["name"],
        "ID": i,
        "Roster Position": r.get("slot", "WR/FLEX"),
        "Salary": r.get("salary", 5500),
        "Game Info": r.get("game", "LV@HOU 08/20/2026 08:00PM ET"),
        "TeamAbbrev": r["team"],
        "AvgPointsPerGame": 0,
    } for i, r in enumerate(rows)])


@pytest.fixture
def dk_file(tmp_path):
    frame = _dk([
        {"name": "Alpha Back", "team": "HOU", "slot": "RB/FLEX"},
        {"name": "Bravo Wide", "team": "LV"},
        {"name": "Charlie Tight", "team": "LV", "slot": "TE/FLEX"},
        {"name": "Delta Arm", "team": "HOU", "slot": "QB"},
        {"name": "Echo Defense", "team": "HOU", "slot": "DST"},
    ])
    path = tmp_path / "DKSalaries.csv"
    frame.to_csv(path, index=False)
    return str(path)


class TestSlateBasics:
    def test_dsts_are_not_skill_players(self, dk_file):
        board = slate.build_slate(dk_file, root="nonexistent")
        assert "Echo Defense" not in set(board["Name"])
        assert set(board["pos"]) <= set(slate.SKILL_POSITIONS)

    def test_roster_slots_fold_to_positions(self, dk_file):
        board = slate.build_slate(dk_file, root="nonexistent")
        assert set(board["pos"]) == {"RB", "WR", "TE", "QB"}

    def test_opponents_come_from_the_game_info_string(self, dk_file):
        board = slate.build_slate(dk_file, root="nonexistent")
        by_team = board.groupby("team")["opp"].first().to_dict()
        assert by_team["HOU"] == "LV"
        assert by_team["LV"] == "HOU"

    def test_it_survives_every_source_being_absent(self, dk_file):
        """A slate report is still worth writing with only the DK file."""
        board = slate.build_slate(dk_file, root="nonexistent")
        assert len(board) == 4
        assert set(board["basis"]) == {"none"}


class TestBasis:
    """`basis` records which source answered, and is never inferred from a value."""

    def _board(self, dk_file, **columns):
        board = slate.build_slate(dk_file, root="nonexistent")
        for name, values in columns.items():
            board[name] = values
        board["basis"] = slate._basis(board)
        return board

    def test_pff_wins_when_it_is_present(self, dk_file):
        board = self._board(dk_file, pff_ppg=[12.0, np.nan, np.nan, np.nan],
                            source=["prior_season"] * 4)
        assert list(board["basis"])[0] == "PFF projection"

    def test_it_falls_through_to_our_own_prior(self, dk_file):
        board = self._board(dk_file, pff_ppg=[np.nan] * 4,
                            source=["prior_season", "draft", "replacement", None])
        assert list(board["basis"]) == ["prior season", "draft capital", "replacement", "none"]


class TestProjection:
    def test_pff_is_used_where_it_exists_and_scaled_by_schedule(self, dk_file):
        board = slate.build_slate(dk_file, root="nonexistent")
        board["pff_ppg"] = [10.0, np.nan, np.nan, np.nan]
        board["sos_mult"] = [1.06, 1.0, 1.0, 1.0]
        board["expected_opportunity"] = [99.0, 20.0, np.nan, np.nan]
        out = slate._blended_projection(board)
        # PFF x schedule, not an average with our own number.
        assert out.iloc[0] == pytest.approx(10.6)

    def test_our_own_number_only_fills_where_pff_is_silent(self, dk_file):
        board = slate.build_slate(dk_file, root="nonexistent")
        board["pff_ppg"] = [np.nan] * 4
        board["sos_mult"] = [1.0] * 4
        board["expected_opportunity"] = [20.0, np.nan, np.nan, np.nan]
        out = slate._blended_projection(board)
        assert out.iloc[0] > 0
        assert pd.isna(out.iloc[1])

    def test_a_cold_start_row_does_not_outrank_a_projected_starter(self, dk_file):
        """Opportunity is not points; the conversion is deliberately held low."""
        board = slate.build_slate(dk_file, root="nonexistent")
        board["pff_ppg"] = [16.0, np.nan, np.nan, np.nan]
        board["sos_mult"] = [1.0] * 4
        board["expected_opportunity"] = [np.nan, 35.0, np.nan, np.nan]
        out = slate._blended_projection(board)
        assert out.iloc[0] > out.iloc[1]


class TestHeat:
    def test_high_is_good_by_default_and_invertible(self):
        good = report._heat_rgb(0.95)
        bad = report._heat_rgb(0.05)
        assert good != bad
        assert report._heat_rgb(0.95, invert=True) == bad

    def test_a_missing_percentile_gets_no_fill(self):
        assert report._heat_rgb(np.nan) is None
        assert report._heat_rgb(None) is None

    def test_a_two_sided_residual_is_neutral_at_zero(self):
        neutral = report._diverging_rgb(0.0, 3.0)
        assert neutral == report._HEAT_NEUTRAL
        assert report._diverging_rgb(3.0, 3.0) != report._diverging_rgb(-3.0, 3.0)

    def test_depth_rank_is_inverted_because_one_is_the_starter(self):
        mode, invert = report.HEAT["depth_rank"]
        assert (mode, invert) == ("pct", True)

    def test_sos_is_not_inverted_because_higher_is_easier(self):
        assert report.HEAT["sos_rating"] == ("pct", False)


class TestWorkbook:
    def _board(self):
        return pd.DataFrame({
            "Name": ["Alpha", "Bravo", "Charlie"],
            "pos": ["RB", "WR", "QB"],
            "team": ["HOU", "LV", "HOU"],
            "opp": ["LV", "HOU", "LV"],
            "basis": ["PFF projection", "draft capital", "none"],
            "proj": [16.0, 3.0, np.nan],
            "pff_ppg": [16.0, np.nan, np.nan],
            "sos_rating": [8.0, 3.0, 8.0],
            "sos_mult": [1.03, 0.98, 1.03],
            "depth_rank": [1, 3, 1],
            "xTD_pg": [0.9, np.nan, np.nan],
            "TD_oe": [-2.5, np.nan, np.nan],
            "scheme_fit": [0.8, -0.4, np.nan],
            "expected_opportunity": [20.0, 4.0, np.nan],
            "availability": [1.0, 0.4, np.nan],
            "Salary": [5500, 5500, 5500],
        })

    def test_it_writes_every_tab(self, tmp_path):
        path = report.write_workbook(self._board(), str(tmp_path / "s.xlsx"), "TEST")
        from openpyxl import load_workbook
        book = load_workbook(path)
        assert book.sheetnames == ["Board", "Red Zone", "Matchup", "Cold Start"]

    def test_a_missing_number_stays_blank_rather_than_becoming_zero(self, tmp_path):
        """A 0.00 would claim the player was projected to score nothing."""
        path = report.write_workbook(self._board(), str(tmp_path / "s.xlsx"))
        from openpyxl import load_workbook
        sheet = load_workbook(path)["Board"]
        header = [c.value for c in sheet[2]]
        column = header.index("Proj") + 1
        values = [sheet.cell(row=r, column=column).value for r in range(3, 6)]
        assert None in values
        assert 0 not in values

    def test_every_sheet_fits_the_zoom_budget(self, tmp_path):
        path = report.write_workbook(self._board(), str(tmp_path / "s.xlsx"))
        from openpyxl import load_workbook
        book = load_workbook(path)
        for name in book.sheetnames:
            sheet = book[name]
            total = sum((sheet.column_dimensions[chr(64 + i)].width or 8.43)
                        for i in range(1, sheet.max_column + 1))
            assert total <= report.WIDTH_BUDGET + 1, (name, total)

    def test_an_all_blank_optional_column_does_not_crash_width_fitting(self, tmp_path):
        """`astype(str).replace("nan","")` returns real NaNs under pandas 3.0.

        The widths loop then calls len() on a float. Regression guard for that.
        """
        board = self._board()
        board["xTD_pg"] = np.nan
        board["TD_oe"] = np.nan
        assert report.write_workbook(board, str(tmp_path / "s.xlsx"))


class TestFigures:
    def _board(self):
        rows = []
        for i, (name, team, pos) in enumerate([
                ("Alpha", "HOU", "RB"), ("Bravo", "LV", "WR"), ("Charlie", "HOU", "TE"),
                ("Delta", "LV", "QB"), ("Echo", "HOU", "WR")]):
            rows.append({"Name": name, "team": team, "pos": pos, "opp": "LV",
                         "proj": 15 - i, "basis": "PFF projection",
                         "xTD_pg": 0.8 - i * 0.1, "TD_oe": (-1) ** i * (i + 1) * 0.7,
                         "sos_rating": 3.0 + i, "scheme_fit": (-1) ** i * (i * 0.3 + 0.2),
                         "opp_man_rate": 20.0})
        return pd.DataFrame(rows)

    def test_all_four_render(self, tmp_path):
        made = report_visuals.build_figures(self._board(), str(tmp_path), label="t")
        assert [caption for caption, _ in made] == [
            "Board", "Touchdown equity", "Strength of schedule", "Coverage fit"]
        for _, path in made:
            assert os.path.getsize(path) > 4000

    def test_a_club_keeps_one_colour_across_figures(self):
        """Slot order is fixed, so a team's hue cannot move between figures or games."""
        first = report_visuals.team_colors(["HOU", "LV"])
        again = report_visuals.team_colors(["LV", "HOU"])
        assert first == again
        assert first["HOU"] == report_visuals.TEAM_SLOTS[0]

    def test_coverage_fit_shows_both_tails(self, tmp_path):
        """A zero-ruled chart of only positive bars wastes half its axis."""
        board = self._board()
        path = report_visuals.figure_scheme_fit(board, str(tmp_path / "f.png"), top_n=4)
        assert path is not None
        values = board["scheme_fit"]
        assert (values > 0).any() and (values < 0).any()

    def test_an_empty_board_yields_no_figures(self, tmp_path):
        assert report_visuals.build_figures(pd.DataFrame(), str(tmp_path)) == []
